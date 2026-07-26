import os
try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    print("⚠️ python-pptx 라이브러리가 설치되지 않았습니다. 터미널에서 'pip install python-pptx'를 실행해주세요.")

from clova.vision.image_text_extractor import ImageTextExtractor
from utils.text_heuristics import extract_frequent_terms

# 아이콘/구분선 같은 장식용 이미지는 OCR 대상에서 제외하기 위한 최소 크기 기준(px)
_MIN_OCR_IMAGE_WIDTH = 150
_MIN_OCR_IMAGE_HEIGHT = 100
# HCX-005 비전 입력 제약(가로세로 비율 1:5~5:1)을 벗어나는 길쭉한 이미지도 제외
_MAX_OCR_ASPECT_RATIO = 5.0
# 비전 호출은 슬라이드당 유료 API 호출이다. 한 장에 이미지가 수십 개 박힌 PPT도 있어서
# (실측: 23장짜리에 후보 이미지 97개) 큰 것부터 이만큼만 읽는다.
_MAX_OCR_IMAGES_PER_SLIDE = 3
# 이미 텍스트박스로 내용이 충분히 들어있는 슬라이드는 이미지를 읽어도 얻을 게 별로 없다.
# 비전은 "텍스트가 거의 없는 슬라이드"를 살리는 용도로만 쓴다.
_OCR_SKIP_TEXT_LENGTH = 50

class PptExtractor:
    def __init__(self):
        self.image_text_extractor = ImageTextExtractor()

    def extract_structured_data(self, file_path: str, topic_hint: str = "", outline_hint: str = "") -> dict:
        """
        [업데이트] PPTX 파일 경로를 입력받아 아래와 같이 구조화된 딕셔너리를 반환합니다.
        1. 발표 주제 및 목차/키워드 자동 추출
        2. 슬라이드 번호별 텍스트 완벽 분리
        3. 텍스트박스가 아니라 이미지(캡처/스캔)로만 된 슬라이드는 HCX-005 비전으로 텍스트 추출 시도

        topic_hint / outline_hint: 사용자가 직접 입력한 발표 주제/목차(Figma 유저 플로우 상의 입력값).
        주어지면 이미지 텍스트 인식 시 문맥으로 함께 제공해 정확도를 높인다. 없으면 힌트 없이 읽는다.
        """
        if not os.path.exists(file_path):
            print(f"❌ 파일을 찾을 수 없습니다: {file_path}")
            return {"metadata": {"topic": "", "keywords": []}, "slides": []}

        context_hint = ""
        if topic_hint:
            context_hint += f"발표 주제: {topic_hint}\n"
        if outline_hint:
            context_hint += f"목차/가이드라인: {outline_hint}\n"

        try:
            prs = Presentation(file_path)
            slides_data = []
            front_text_for_analysis = []
            all_slide_texts = []

            for i, slide in enumerate(prs.slides):
                slide_texts = []

                # 1차: 텍스트박스에서 바로 읽히는 글자를 순서대로 모으고,
                #      이미지로만 들어간 내용은 후보로만 쌓아둔다(비전 호출은 유료라 뒤에서 추려서 한다).
                ocr_candidates = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())
                    elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        # 아이콘/구분선 같은 장식용 이미지는 크기/비율로 걸러내고,
                        # 내용이 있을 법한 이미지만 후보로 남긴다.
                        try:
                            image = shape.image
                            width_px, height_px = image.size
                        except Exception:
                            continue
                        if width_px < _MIN_OCR_IMAGE_WIDTH or height_px < _MIN_OCR_IMAGE_HEIGHT:
                            continue
                        aspect_ratio = max(width_px, height_px) / max(1, min(width_px, height_px))
                        if aspect_ratio > _MAX_OCR_ASPECT_RATIO:
                            continue
                        ocr_candidates.append((width_px * height_px, image))

                # 2차: 텍스트가 거의 없는 슬라이드(= 이미지가 곧 내용인 장표)만 비전으로 읽는다.
                #      큰 이미지일수록 본문일 확률이 높으니 넓이 순으로 상위 몇 장만 본다.
                if ocr_candidates and len("".join(slide_texts)) < _OCR_SKIP_TEXT_LENGTH:
                    ocr_candidates.sort(key=lambda item: item[0], reverse=True)
                    for _, image in ocr_candidates[:_MAX_OCR_IMAGES_PER_SLIDE]:
                        ocr_text = self.image_text_extractor.extract_text_from_image(
                            image.blob, context_hint, image.content_type
                        )
                        if ocr_text.strip():
                            slide_texts.append(ocr_text.strip())

                slide_content = "\n".join(slide_texts)

                # [요구사항 5 반영] 슬라이드별 텍스트가 존재하는 경우에만 객체로 분리하여 추가
                if slide_content.strip():
                    slides_data.append({
                        "slide_number": i + 1,
                        "content": slide_content
                    })

                    all_slide_texts.extend(slide_texts)

                    # 주제 및 목차 탐지를 위해 초반 1~3장 텍스트는 별도로도 수집
                    if i < 3:
                        front_text_for_analysis.extend(slide_texts)

            # [요구사항 4 반영] 초반 슬라이드 텍스트 기반 주제 추출 + 전체 슬라이드 기반 키워드 추출
            topic, keywords = self._extract_metadata(front_text_for_analysis, all_slide_texts)

            print(f"✅ PPT 구조화 추출 완료! (총 {len(slides_data)}장 분석)")
            return {
                "metadata": {
                    "topic": topic,
                    "keywords": keywords
                },
                "slides": slides_data
            }

        except Exception as e:
            print(f"❌ PPT 추출 중 오류 발생: {e}")
            return {"metadata": {"topic": "", "keywords": []}, "slides": []}

    def _extract_metadata(self, front_texts: list, all_texts: list = None):
        """초반 슬라이드 텍스트로 발표 주제를, 전체 슬라이드 텍스트로 키워드를 휴리스틱하게 추론합니다."""
        if not front_texts:
            return "주제 미상", []

        all_texts = all_texts if all_texts else front_texts

        # 보통 첫 번째 텍스트 덩어리가 발표 주제(Title)일 확률이 높음
        topic = front_texts[0].strip()
        keywords = []

        # '목차', 'index', 'agenda' 등의 단어가 포함된 슬라이드가 있다면 그 직후 텍스트를 키워드로 간주
        for i, text in enumerate(all_texts):
            lower_text = text.lower()
            if any(keyword in lower_text for keyword in ['목차', 'index', 'contents', '순서', 'agenda', '차례', 'outline']):
                keywords = all_texts[i + 1: i + 6]
                break

        # 목차 슬라이드가 없다면, 전체 슬라이드 텍스트에서 자주 등장하는 단어를 키워드로 추론
        if not keywords:
            keywords = self._extract_keywords_by_frequency(all_texts, topic)

        return topic, keywords

    def _extract_keywords_by_frequency(self, texts: list, topic: str, top_n: int = 5) -> list:
        """목차 슬라이드가 없는 PPT를 위한 대체 키워드 추출: 전체 텍스트에서 자주 등장하는 단어를 뽑는다."""
        return extract_frequent_terms(texts, exclude={topic}, top_n=top_n)

# ==========================================
# 🧪 [테스트 코드]
# ==========================================
if __name__ == "__main__":
    extractor = PptExtractor()
    test_file = "sample.pptx" 
    
    if os.path.exists(test_file):
        # 기존 통글 추출 대신 구조화된 데이터 추출 메서드 사용
        result_data = extractor.extract_structured_data(test_file)
        
        import json
        print("\n✨ [추출된 PPT 구조화 데이터] ✨")
        print(json.dumps(result_data, indent=2, ensure_ascii=False))
    else:
        print(f"⚠️ '{test_file}' 파일이 존재하지 않아 추출 테스트를 건너뜁니다.")